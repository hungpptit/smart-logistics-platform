import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# Load grouped provinces & wards data exported from DB
with open('d:/smart-logistics-platform/backend/provinces_wards_grouped.json', 'r', encoding='utf-8') as f:
    grouped_data = json.load(f)

wb = openpyxl.Workbook()

# Sheet 1: Main Order Template
ws_main = wb.active
ws_main.title = "Đơn Hàng Loạt Case 1 (10 Đơn)"
ws_main.views.sheetView[0].showGridLines = True

# Sheet 2: Reference Administrative Lookup Sheet
ws_lookup = wb.create_sheet(title="DanhMucHanhChinh")
ws_lookup.views.sheetView[0].showGridLines = True

# 1. Write Column A in Lookup Sheet: All Provinces
ws_lookup.cell(row=1, column=1, value="Tỉnh/Thành Phố").font = Font(bold=True, color="1E3A8A")

prov_names = list(grouped_data.keys())
for idx, prov_fullName in enumerate(prov_names, start=2):
    ws_lookup.cell(row=idx, column=1, value=prov_fullName)

ws_lookup.column_dimensions['A'].width = 30

def make_valid_range_name(name):
    return name.replace(" ", "_").replace(".", "").replace("-", "_")

# 2. Write Wards for each Province in separate columns
for prov_idx, prov_fullName in enumerate(prov_names, start=2):
    col_letter = get_column_letter(prov_idx)
    range_name = make_valid_range_name(prov_fullName)
    
    ws_lookup.cell(row=1, column=prov_idx, value=prov_fullName).font = Font(bold=True)
    
    wards_list = grouped_data[prov_fullName].get('wards', [])
    for w_idx, ward_fullName in enumerate(wards_list, start=2):
        ws_lookup.cell(row=w_idx, column=prov_idx, value=ward_fullName)
    
    ws_lookup.column_dimensions[col_letter].width = 32
    
    last_row = max(len(wards_list) + 1, 2)
    formula_ref = f"DanhMucHanhChinh!${col_letter}$2:${col_letter}${last_row}"
    
    new_defined_name = openpyxl.workbook.defined_name.DefinedName(range_name, attr_text=formula_ref)
    wb.defined_names.add(new_defined_name)

# --- MAIN SHEET HEADERS (26 Standard Columns) ---
headers = [
    # STEP 1: NGƯỜI GỬI (KHO 97 MAN THIỆN - TĂNG NHƠN PHÚ)
    ("HoTenNguoiGui", "1. Họ tên người gửi", "1E3A8A"),
    ("SoDienThoaiGui", "1. SĐT người gửi", "1E3A8A"),
    ("DiaChiGui", "1. Địa chỉ lấy hàng (Số nhà, đường)", "1E3A8A"),
    ("PhuongXaGui", "1. Phường/Xã gửi (Dropdown theo Tỉnh)", "1E3A8A"),
    ("TinhThanhGui", "1. Tỉnh/TP gửi (Dropdown)", "1E3A8A"),

    # STEP 2: NGƯỜI NHẬN (KHU VỰC LINH TRUNG)
    ("HoTenNguoiNhan", "2. Họ tên người nhận*", "4338CA"),
    ("SoDienThoaiNhan", "2. SĐT người nhận*", "4338CA"),
    ("DiaChiGiaoHang", "2. Địa chỉ giao (Số nhà, đường)*", "4338CA"),
    ("PhuongXaNhan", "2. Phường/Xã nhận (Dropdown theo Tỉnh)*", "4338CA"),
    ("TinhThanhNhan", "2. Tỉnh/TP nhận (Dropdown)*", "4338CA"),

    # STEP 3: GÓI HÀNG
    ("TenHangHoa", "3. Mô tả danh mục hàng hóa*", "D97706"),
    ("TrongLuongKg", "3. Trọng lượng (kg)*", "D97706"),
    ("DaiCm", "3. Dài (cm)", "D97706"),
    ("RongCm", "3. Rộng (cm)", "D97706"),
    ("CaoCm", "3. Cao (cm)", "D97706"),
    ("GiaTriKhaiGia", "3. Khai giá hàng hóa (VNĐ)", "D97706"),
    ("HangDeVo", "3. Hàng dễ vỡ?", "D97706"),
    ("YeuCauNhietDo", "3. Yêu cầu nhiệt độ", "D97706"),

    # STEP 4: DỊCH VỤ & THANH TOÁN
    ("GoiCuocDichVu", "4. Gói cước dịch vụ", "059669"),
    ("HinhThucGuiHang", "4. Hình thức gửi hàng", "059669"),
    ("NgayLayHang", "4. Ngày lấy hàng", "059669"),
    ("CaThuHang", "4. Ca thu hàng", "059669"),
    ("NguoiChiuPhi", "4. Người chịu phí", "059669"),
    ("HinhThucThanhToan", "4. Hình thức thanh toán", "059669"),
    ("TienThuHoCOD", "4. Tiền thu hộ COD (VNĐ)", "059669"),
    ("GhiChuGiaoHang", "4. Ghi chú dặn Shipper", "059669"),
]

font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

ws_main.row_dimensions[1].height = 40

for col_idx, (code, title, fill_color) in enumerate(headers, 1):
    cell = ws_main.cell(row=1, column=col_idx, value=title)
    cell.font = font_header
    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    cell.alignment = align_center
    cell.border = thin_border

# 10 Real-world Orders in Linh Trung Area
case1_orders_data = [
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Nguyễn Thị Mai Anh", "0912111001", "120 Đường Linh Trung", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Bộ mỹ phẩm dưỡng da Hàn Quốc Innisfree", 0.8, 15, 10, 8, 850000, "Có", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 350000, "Cho khách xem hàng trước khi nhận"
    ],
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Trần Đình Trọng", "0912111002", "45 Đường Lê Văn Chí", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Tai nghe Bluetooth chụp tai chống ồn Sony", 0.5, 20, 15, 10, 1200000, "Có", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 450000, "Gọi điện trước khi giao 15 phút"
    ],
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Lê Hoàng Nam", "0912111003", "78 Đường Hoàng Diệu 2", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Giày thể thao sneaker nam cao cấp size 42", 1.2, 30, 20, 12, 950000, "Không", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 620000, "Giao giờ hành chính tại văn phòng"
    ],
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Phạm Ngọc Hân", "0912111004", "15 Đường Số 5, Khu phố 2", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Set đầm xòe công sở thiết kế hoa nhí", 0.6, 25, 20, 5, 550000, "Không", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 280000, "Giao buổi sáng trước 11h"
    ],
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Vũ Minh Tuấn", "0912111005", "210 Quốc Lộ 1K", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Bàn phím cơ Bluetooth RGB Gaming 87 phím", 1.5, 35, 15, 6, 1500000, "Có", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 890000, "Kiện hàng điện tử, giao cẩn thận"
    ],
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Đặng Thu Trang", "0912111006", "32 Đường Số 7, Linh Trung", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Nồi chiên không dầu điện tử Tefal 5.5L", 3.5, 35, 35, 38, 2200000, "Có", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 1150000, "Hàng kích thước lớn, giao tận cửa nhà"
    ],
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Hoàng Quốc Việt", "0912111007", "88 Đường Kha Vạn Cân", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Bộ dụng cụ cơ khí sửa chữa đa năng 108 món", 2.2, 38, 28, 9, 750000, "Không", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 380000, "Nhà trong hẻm, gọi khi đến đầu đường"
    ],
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Bùi Thị Thu Thảo", "0912111008", "102 Đường Linh Trung", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Đèn bàn LED bảo vệ thị lực có sạc không dây", 1.1, 30, 18, 12, 680000, "Có", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 520000, "Giao trong ngày hôm nay"
    ],
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Ngô Thành Đạt", "0912111009", "15 Đường Võ Văn Ngân", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Áo khoác gió thể thao 2 lớp chống nước", 0.7, 25, 20, 5, 650000, "Không", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 490000, "Giao tại cửa hàng điện thoại"
    ],
    [
        "PTIT HCM 1", "0987654321", "97 Man Thiện", "Phường Tăng Nhơn Phú", "Thành phố Hồ Chí Minh",
        "Dương Kim Ngân", "0912111010", "60 Đường Chương Dương", "Phường Linh Xuân", "Thành phố Hồ Chí Minh",
        "Bình giữ nhiệt Lock&Lock dung tích 800ml", 0.5, 22, 10, 10, 380000, "Không", "Thường",
        "STANDARD (Tiêu chuẩn)", "PICKUP (Shipper lấy tận nơi)", "Hôm nay", "Ca Sáng (08:00 - 12:00)", "SENDER (Người gửi trả)", "COD (Thu hộ COD)", 230000, "Gọi trước khi đến 10 phút"
    ]
]

align_data = Alignment(horizontal="left", vertical="center")
font_data = Font(name="Arial", size=9)

for r_idx, row_values in enumerate(case1_orders_data, start=2):
    ws_main.row_dimensions[r_idx].height = 24
    for c_idx, val in enumerate(row_values, start=1):
        cell = ws_main.cell(row=r_idx, column=c_idx, value=val)
        cell.font = font_data
        cell.alignment = align_data
        cell.border = thin_border

# --- DATA VALIDATIONS ---
prov_range = f"=DanhMucHanhChinh!$A$2:$A${len(prov_names)+1}"

dv_prov_sender = DataValidation(type="list", formula1=prov_range, allow_blank=True)
dv_prov_receiver = DataValidation(type="list", formula1=prov_range, allow_blank=True)
dv_ward_sender = DataValidation(type="list", formula1='=INDIRECT(SUBSTITUTE(SUBSTITUTE(E2, " ", "_"), ".", ""))', allow_blank=True)
dv_ward_receiver = DataValidation(type="list", formula1='=INDIRECT(SUBSTITUTE(SUBSTITUTE(J2, " ", "_"), ".", ""))', allow_blank=True)

dv_fragile = DataValidation(type="list", formula1='"Có,Không"', allow_blank=True)
dv_service = DataValidation(type="list", formula1='"STANDARD (Tiêu chuẩn),EXPRESS (Hỏa tốc 2h)"', allow_blank=True)
dv_pickup_type = DataValidation(type="list", formula1='"PICKUP (Shipper lấy tận nơi),DROP_OFF (Gửi tại bưu cục)"', allow_blank=True)
dv_pickup_date = DataValidation(type="list", formula1='"Hôm nay,Ngày mai"', allow_blank=True)
dv_pickup_shift = DataValidation(type="list", formula1='"Ca Sáng (08:00 - 12:00),Ca Chiều (13:00 - 18:00)"', allow_blank=True)
dv_fee_payer = DataValidation(type="list", formula1='"SENDER (Người gửi trả),RECEIVER (Người nhận trả)"', allow_blank=True)
dv_payment = DataValidation(type="list", formula1='"CASH (Tiền mặt),COD (Thu hộ COD),BANK_TRANSFER (Chuyển khoản)"', allow_blank=True)

ws_main.add_data_validation(dv_prov_sender)
ws_main.add_data_validation(dv_ward_sender)
ws_main.add_data_validation(dv_prov_receiver)
ws_main.add_data_validation(dv_ward_receiver)
ws_main.add_data_validation(dv_fragile)
ws_main.add_data_validation(dv_service)
ws_main.add_data_validation(dv_pickup_type)
ws_main.add_data_validation(dv_pickup_date)
ws_main.add_data_validation(dv_pickup_shift)
ws_main.add_data_validation(dv_fee_payer)
ws_main.add_data_validation(dv_payment)

dv_prov_sender.add("E2:E50")
dv_ward_sender.add("D2:D50")
dv_prov_receiver.add("J2:J50")
dv_ward_receiver.add("I2:I50")
dv_fragile.add("Q2:Q50")
dv_service.add("S2:S50")
dv_pickup_type.add("T2:T50")
dv_pickup_date.add("U2:U50")
dv_pickup_shift.add("V2:V50")
dv_fee_payer.add("W2:W50")
dv_payment.add("X2:X50")

# Auto-fit column widths
for col in ws_main.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        if len(val_str) > max_len:
            max_len = len(val_str)
    ws_main.column_dimensions[col_letter].width = max(max_len + 3, 18)

excel_filepath = "d:/smart-logistics-platform/DATA_TEST_CASE1_10_DON_HANG.xlsx"
wb.save(excel_filepath)
print(f"Successfully generated 10-order Case 1 Excel template: {excel_filepath}")
